import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Optional

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class Book(BaseModel):
    isbn: str
    title: str
    author: str
    publisher: str
    cost: float

books_db: List[Book] = []

def load_initial_books():
    raw_data = [
        [2859, "1 to 9 Konkan Rly Photographs incl. Bridges & Tunnels etc.", "-", "-", 0.0],
        [2860, "Bombay A to Z Guide", "Mrs. S.M. Karmarkar Bombay", "-", 11.0],
        [2861, "Pune City (Illustrated) Information & Map", "-", "-", 30.0],
        [2862, "Bombay Tourist Map Folder", "-", "MTDC", 20.0],
        [2863, "Aurangabad Tourist Map & Folder", "-", "MTDC", 0.0],
        [2864, "Bombay Tourist Map Folder", "-", "MTDC", 0.0],
        [2865, "Maharashtra Places 1 to ...", "-", "-", 0.0],
        [2868, "Bombay Tourist Map", "-", "MTDC", 0.0],
        [2869, "Maharashtra State Road Map", "Tej Bros.", "Tej Bros.", 4.0],
        [2870, "Aurangabad Ellora Ajanta Paithan Folder", "-", "MTDC", 0.0],
        [16, "Vayu Purana", "P. N. Joshi", "Prasad Prakashan", 0.0],
        [17, "Hiranyagarbha - Kahani Brahmandnayakachi", "Sakharam Amle", "Rasik Aantabharti", 400.0],
        [18, "Shrikhetra Neera Narsinhpur", "Vasant Aragade", "-", 0.0],
        [19, "Shri Ganesh Pratap", "Madhav Parchure", "Indira Printing Press", 15.0],
        [20, "Samarth Vangmay Jayanti", "Ch. P. Bhishikar", "Prasad Prakashan", 0.0],
        [21, "Purva-Paschim", "Shanta Paranjape", "Morya Prakashan", 20.0],
        [23, "Sangeet Vichar", "Ashok Ranade", "Popular Prakashan", 0.0],
        [24, "Shri Vidarbha Ganesh Darshan Kosh", "Madhukarrao Ashte", "-", 0.0],
        [1, "Samarthanchi Lokgeete", "Nandkumar Marathe", "Prasad Prakashan", 500.0],
        [2, "Samarthanchi Updeshpar", "-", "-", 550.0],
        [3, "Samarth Vangmay Jayanti", "Dr. Indu Limaye", "-", 350.0],
        [4, "Yogi Samarthyacha", "Meghshyam Savakar", "-", 400.0],
        [5, "Samarthanchi Prasadik Vani", "Ahilya Deshmukh", "-", 550.0],
        [6, "Shivram", "Dadasaheb Rajshirke", "Shri Vimal Publishers", 250.0],
        [7, "Yuddha Pratapgadche", "Major Mukund Joshi", "Snehvardhan Prakashan", 120.0],
        [8, "Bhasakrut Charudatta Shudrakkrut Mrichchhakatika", "Ram Shewalkar", "Prasad Prakashan", 250.0],
        [9, "Kshatriya Kulacha Itihas Aitihasik Kagadpatre", "Sunil Hapurkar", "Gagangiri Prakashan", 0.0],
        [10, "Shri Ganesh Geeta (Marathi)", "Pradip Rasne", "Latika Prakashan", 700.0],
        [2842, "Tourist map Folder Pune - The Deccan Queen", "-", "MTDC", 10.0],
        [2843, "Tourist Guide map of Jodhpur", "-", "IMS Jodhpur", 6.2],
        [2848, "Discover Maharashtra Bhimashankar MTDC", "-", "-", 5.0],
        [2849, "Dajipur Radhanagari, Rehekuri - Nanaj...", "-", "-", 0.0],
        [2850, "do - The Melghat Chikhaldara Adventure", "-", "-", 0.0],
        [2851, "do - The Navegaon, Nagzira, Tadoba", "-", "-", 0.0],
        [2852, "Pune Guide", "B.H. Joshi", "Swayam P.", 40.0],
        [2854, "Trekking East - leaflet", "Miss P. Barve", "Pune", 2.0],
        [2830, "Road Map of Delhi with D.T.C Bus Routes", "-", "India Book Depot House", 13.45],
        [2831, "Latest Tourist Guide Goa - with Road Map", "Asha Pub House", "Delhi", 8.0],
        [2832, "Map of Haryana (New & Latest)", "-", "-", 3.0],
        [2833, "States of Our Union - 8 Haryana - with Map", "Shashi Bhushan", "Shiksha Chatra Bhandar", 2.0],
        [2834, "Amritsar - Spiritual Centre of the Punjab", "-", "Govt. of Punjab", 2.0],
        [2836, "Road Map of Haryana", "Burjor M. D'Mehta", "Indian Book Depot", 2.0],
        [2838, "States of Our Union - Punjab", "Sher Singh Gupta", "Govt. of India", 2.0],
        [31, "Namsangeeta (Akaravya Geeta)", "Ki. M. Shirke", "Hrishikesh Prakashan", 50.0],
        [32, "Ekadash", "Gopal Dandekar", "-", 0.0],
        [33, "Nisarga Ramya Matheran", "-", "-", 0.0],
        [34, "Stotra - Shrisukta", "-", "Anmol Prakashan", 40.0],
        [35, "Dnyaneshwaritil Granthisamvad", "Shripad Dhaygude", "Anmol Prakashan", 280.0],
        [36, "Shri Gajanan Maharaj - Shegaon Vijay", "Arun Gokhale", "Anmol Prakashan", 70.0],
        [38, "Samarth Vangmay Jayanti", "Indu Limaye", "Prasad Prakashan", 0.0],
        [43, "Shri Samarth Granth Bhandar", "H. B. P. Pangarkar", "Jyoti Dhawale", 225.0],
        [44, "Jyotibache Darshan", "Maruti Chitampalli", "A. L. Kulkarni", 0.0],
        [2817, "Our Feathered Friends", "V. C. Chopra", "ACBT New Delhi", 0.0],
        [2819, "Goa - Travel Guide with Multicoloured Map", "K. C. Damodaran", "Streamlines Publn Cochin", 7.5],
        [2820, "Tirupati - A.P.", "Govt. of A.P.", "-", 5.0],
        [2821, "A.P. Tourist Map", "\"", "-", 8.0],
        [2822, "Goa (with Panaji) (Discover India series)", "TTK Scale 1:160,000", "-", 90.0],
        [2824, "Tourist Guide to Goa (The land of Beaches & Churches)", "-", "Sura Madras", 8.0],
        [2825, "Temples of Tirumala, Tirupati & Tiruchanur", "Dr M. Rama Rao", "T.T. Devasthanam", 3.0],
        [2826, "Tirupati - Tirumala", "-", "Lotus Pub House", 3.0],
        [2827, "Tamil Nadu - a land of splendour and serenity", "-", "The Director of Tourism", 0.0]
    ]
    
    for row in raw_data:
        books_db.append(Book(
            isbn=str(row[0]),
            title=str(row[1]),
            author=str(row[2]),
            publisher=str(row[3]),
            cost=float(row[4])
        ))

load_initial_books()

@app.get("/books", response_model=List[Book])
def get_books(search_by: Optional[str] = None, query: Optional[str] = None):
    if query and search_by:
        q = query.lower()
        if search_by == "title":
            return [b for b in books_db if q in b.title.lower()]
        elif search_by == "author":
            return [b for b in books_db if q in b.author.lower()]
        elif search_by == "publisher":
            return [b for b in books_db if q in b.publisher.lower()]
        elif search_by == "isbn":
            return [b for b in books_db if q in b.isbn.lower()]
    return books_db

@app.post("/books", response_model=Book)
def add_book(book: Book):
    if any(b.isbn == book.isbn for b in books_db):
        raise HTTPException(status_code=400, detail="Book with this ISBN already exists")
    books_db.append(book)
    return book

@app.put("/books/{isbn}", response_model=Book)
def update_book(isbn: str, updated_book: Book):
    for index, book in enumerate(books_db):
        if book.isbn == isbn:
            books_db[index] = updated_book
            return updated_book
    raise HTTPException(status_code=404, detail="Book not found")

@app.delete("/books/{isbn}")
def delete_book(isbn: str):
    global books_db
    for index, book in enumerate(books_db):
        if book.isbn == isbn:
            deleted_book = books_db.pop(index)
            return {"message": f"Book '{deleted_book.title}' deleted successfully"}
    raise HTTPException(status_code=404, detail="Book not found")

@app.get("/export-excel")
def export_excel():
    data = []
    for b in books_db:
        data.append([b.isbn, b.title, b.author, b.publisher, b.cost])
        
    df = pd.DataFrame(data, columns=["Entry No.", "Name of the Book", "Writer", "Publisher", "Price"])
    
    file_path = "Book_Accession_Register.xlsx"
    writer = pd.ExcelWriter(file_path, engine='openpyxl')
    df.to_excel(writer, index=False, sheet_name='Books Data')
    
    worksheet = writer.sheets['Books Data']
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    border_style = Side(border_style="thin", color="000000")
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    for col in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    column_widths = {'A': 10, 'B': 45, 'C': 25, 'D': 25, 'E': 10}
    for i, col_letter in enumerate(column_widths.keys(), start=1):
        worksheet.column_dimensions[col_letter].width = column_widths[col_letter]
        for row in range(2, len(df) + 2):
            cell = worksheet.cell(row=row, column=i)
            cell.border = border
            cell.alignment = alignment
            
    writer.close()
    return {"message": f"File generated successfully: {file_path}"}